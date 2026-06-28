import os
import shutil
import glob
import re
import traceback
import openpyxl
import pdfplumber
import pytesseract
from pdf2image import convert_from_path

def log_error(base_dir, message):
    log_path = os.path.join(base_dir, "error.log")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(message + "\n")

def perform_ocr(pdf_path):
    """Fallback OCR using pytesseract if the PDF is scanned."""
    try:
        images = convert_from_path(pdf_path)
        text = ""
        for img in images:
            text += pytesseract.image_to_string(img, lang='spa') + "\n"
        return text
    except Exception as e:
        return f"OCR Failed: {str(e)}"

def extract_text_from_pdf(pdf_path):
    text = ""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        pass
    
    # If no text found via pdfplumber (scanned document), use OCR fallback
    if not text.strip():
        text = perform_ocr(pdf_path)
        
    return text

def process_inmobiliaria(year, data_dir, results_dir):
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    target_filename = f"{year} Agua y luz.xlsx"
    target_path = os.path.join(results_dir, target_filename)
    
    # 1. Setup Excel template
    if not os.path.exists(target_path):
        ref_path = os.path.join(base_dir, "references", "2025 Agua y luz.xlsx")
        if not os.path.exists(ref_path):
            log_error(base_dir, f"Missing template file: {ref_path}")
            raise FileNotFoundError(f"Missing template file: {ref_path}")
        
        os.makedirs(results_dir, exist_ok=True)
        shutil.copy(ref_path, target_path)
        
        # Empty existing values in the copied template
        wb = openpyxl.load_workbook(target_path)
        
        if 'Luz' in wb.sheetnames:
            ws_luz = wb['Luz']
            # Delete rows 2 to 13
            ws_luz.delete_rows(2, 12)
            
        if 'Agua' in wb.sheetnames:
            ws_agua = wb['Agua']
            # Delete rows 2 to 7
            ws_agua.delete_rows(2, 6)
            
        wb.save(target_path)
        wb.close()
    
    wb = openpyxl.load_workbook(target_path)
    
    # --- PROCESS AGUA ---
    agua_pdfs = glob.glob(os.path.join(data_dir, "Comunidad Las Colinas Cuota*.pdf"))
    if 'Agua' in wb.sheetnames:
        ws_agua = wb['Agua']
        
        # Find next empty row
        agua_next_row = ws_agua.max_row + 1
        
        for pdf_path in agua_pdfs:
            text = extract_text_from_pdf(pdf_path)
            
            # Search for Concepto: ... AGUA
            # E.g. Concepto: CUOTA... AGUA (lect. 02/12/2025) 4352-4332:20m3 :38,8
            lines = text.split('\n')
            found_agua = False
            for line in lines:
                if 'Concepto:' in line and 'AGUA' in line:
                    found_agua = True
                    try:
                        # Extract Recibo: e.g. (lect. 02/12/2025) -> 12-2025
                        recibo_match = re.search(r'\(lect\.\s*\d{2}/(\d{2})/(\d{4})\)', line)
                        if recibo_match:
                            recibo = f"{recibo_match.group(1)}-{recibo_match.group(2)}"
                        else:
                            recibo = f"Unknown-{year}"
                            
                        # Extract Vol: e.g. 4352-4332 -> 4352
                        vol_match = re.search(r'\)\s*(\d+)-\d+', line)
                        vol = vol_match.group(1) if vol_match else "0"
                        
                        # Extract Importe: e.g. :38,8 at the end
                        importe_match = re.search(r':\s*(\d+[,.]\d+)$', line.strip())
                        if importe_match:
                            importe_str = importe_match.group(1).replace(',', '.')
                            importe = float(importe_str)
                        else:
                            importe = 0.0
                            
                        # Clear old 'Nuevo recibo'
                        for r in range(2, agua_next_row):
                            if ws_agua.cell(row=r, column=1).value == "Nuevo recibo":
                                ws_agua.cell(row=r, column=1).value = ""
                                
                        # Write to sheet (Status, Recibo, Vol, Importe)
                        ws_agua.cell(row=agua_next_row, column=1).value = "Nuevo recibo"
                        ws_agua.cell(row=agua_next_row, column=2).value = recibo
                        ws_agua.cell(row=agua_next_row, column=3).value = int(vol)
                        ws_agua.cell(row=agua_next_row, column=4).value = importe
                        
                        agua_next_row += 1
                    except Exception as e:
                        log_error(base_dir, f"Failed to parse AGUA data in {pdf_path}: {e}\nLine: {line}")
            
            if not found_agua:
                # Normal for odd months, just skip
                pass

    # --- PROCESS LUZ ---
    luz_pdfs = glob.glob(os.path.join(data_dir, "Electricidad*.pdf"))
    if 'Luz' in wb.sheetnames:
        ws_luz = wb['Luz']
        luz_next_row = ws_luz.max_row + 1
        
        for pdf_path in luz_pdfs:
            # Extract month/year from filename: Electricidad XX-202Y.pdf
            filename = os.path.basename(pdf_path)
            recibo_match = re.search(r'(\d{2}-\d{4})', filename)
            recibo = recibo_match.group(1) if recibo_match else f"Unknown-{year}"
            
            text = extract_text_from_pdf(pdf_path)
            
            # Find IMPORTE TOTAL: EUROS   *** 51,98
            importe = 0.0
            found_importe = False
            lines = text.split('\n')
            for line in lines:
                if 'IMPORTE TOTAL: EUROS' in line:
                    importe_match = re.search(r'\*{1,}\s*(\d+[,.]\d+)', line)
                    if importe_match:
                        importe_str = importe_match.group(1).replace(',', '.')
                        importe = float(importe_str)
                        found_importe = True
                        break
                        
            if not found_importe:
                # Try to find an alternative layout
                for line in lines:
                    if 'IMPORTE TOTAL' in line:
                        importe_match = re.search(r'(\d+[,.]\d+)$', line.strip())
                        if importe_match:
                            importe_str = importe_match.group(1).replace(',', '.')
                            importe = float(importe_str)
                            found_importe = True
                            break
                            
            if not found_importe:
                 log_error(base_dir, f"Failed to find IMPORTE TOTAL in {pdf_path}")
                 
            # Clear old 'Nuevo recibo'
            for r in range(2, luz_next_row):
                if ws_luz.cell(row=r, column=1).value == "Nuevo recibo":
                    ws_luz.cell(row=r, column=1).value = ""
                    
            # Write to sheet (Status, Recibo, Importe)
            ws_luz.cell(row=luz_next_row, column=1).value = "Nuevo recibo"
            ws_luz.cell(row=luz_next_row, column=2).value = recibo
            ws_luz.cell(row=luz_next_row, column=3).value = importe
            
            luz_next_row += 1

    wb.save(target_path)
    wb.close()
    
    return target_path
